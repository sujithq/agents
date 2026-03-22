#!/usr/bin/env python3
"""
Extract speaker notes from PPTX files.

This script extracts speaker notes from each slide in a PowerPoint presentation
and saves them as individual text files.
"""

import argparse
import sys
from pathlib import Path


def extract_notes(pptx_path, output_dir):
    """
    Extract speaker notes from PPTX file.

    Args:
        pptx_path: Path to input PPTX file
        output_dir: Directory to save notes text files
    """
    try:
        from pptx import Presentation
    except ImportError:
        print("Error: python-pptx library required", file=sys.stderr)
        print("Install with: pip install python-pptx", file=sys.stderr)
        sys.exit(1)

    pptx_path = Path(pptx_path).resolve()
    output_dir = Path(output_dir).resolve()

    # Validate input
    if not pptx_path.exists():
        print(f"Error: PPTX file not found: {pptx_path}", file=sys.stderr)
        sys.exit(1)

    # Create output directory
    output_dir.mkdir(parents=True, exist_ok=True)

    print(f"Extracting notes from: {pptx_path}")
    print(f"Output directory: {output_dir}")

    # Load presentation
    try:
        prs = Presentation(str(pptx_path))
    except Exception as e:
        print(f"Error loading PPTX file: {e}", file=sys.stderr)
        sys.exit(1)

    # Extract notes from each slide
    total_slides = len(prs.slides)
    slides_with_notes = 0

    for i, slide in enumerate(prs.slides, start=1):
        output_path = output_dir / f"slide-{i:03d}.txt"

        # Extract notes text
        notes_text = ""
        try:
            if slide.has_notes_slide:
                notes_slide = slide.notes_slide
                if notes_slide and notes_slide.notes_text_frame:
                    notes_text = notes_slide.notes_text_frame.text.strip()
        except Exception as e:
            print(f"Warning: Error extracting notes from slide {i}: {e}", file=sys.stderr)

        # Write notes to file (even if empty)
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(notes_text)

        if notes_text:
            slides_with_notes += 1
            word_count = len(notes_text.split())
            print(f"  - slide-{i:03d}.txt ({word_count} words)")
        else:
            print(f"  - slide-{i:03d}.txt (no notes)")

    print(f"\nExtracted notes from {total_slides} slides")
    print(f"  - {slides_with_notes} slides with notes")
    print(f"  - {total_slides - slides_with_notes} slides without notes")

    return total_slides


def main():
    """Main entry point."""
    parser = argparse.ArgumentParser(
        description="Extract speaker notes from PPTX files"
    )
    parser.add_argument(
        "--pptx",
        required=True,
        help="Path to input PPTX file"
    )
    parser.add_argument(
        "--out",
        required=True,
        help="Output directory for notes text files"
    )

    args = parser.parse_args()

    extract_notes(args.pptx, args.out)


if __name__ == "__main__":
    main()
